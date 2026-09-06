"""Excel workbook: one tab per analysis, formatted as filterable tables with frozen headers and hyperlinks."""
from __future__ import annotations
import json, re
import pandas as pd
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
def _san(v):
    return _ILLEGAL.sub("", v) if isinstance(v, str) else v
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles import Font, Alignment

MAX_ROWS = 200000

def _write(ws, df: pd.DataFrame, table_name: str, link_cols=(), widths=None, max_width=60):
    df = df.loc[:, ~df.columns.duplicated()].copy()
    for c in df.columns:
        if df[c].dtype == object:
            df[c] = df[c].map(lambda v: "|".join(v) if isinstance(v, list) else v)
    df = df.where(pd.notna(df), None)
    ws.append([str(c) for c in df.columns])
    for row in df.itertuples(index=False):
        ws.append([_san(None if (isinstance(v, float) and pd.isna(v)) else (v.isoformat() if hasattr(v, "isoformat") else v)) for v in row])
    n, m = len(df) + 1, len(df.columns)
    if m == 0: return
    ref = f"A1:{get_column_letter(m)}{max(n, 2)}"
    t = Table(displayName=table_name[:250].replace(" ", "_").replace("-", "_"), ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(t); ws.freeze_panes = "A2"
    for j, c in enumerate(df.columns, 1):
        if c in link_cols:
            for i in range(2, n + 1):
                cell = ws.cell(row=i, column=j)
                if cell.value and str(cell.value).startswith("http"):
                    cell.hyperlink = str(cell.value); cell.font = Font(color="0563C1", underline="single")
        sample = max((len(str(v)) for v in df[c].head(200) if v is not None), default=10)
        ws.column_dimensions[get_column_letter(j)].width = min(max(10, int(sample or 10) + 2, len(str(c)) + 2), max_width)

def _text_sheet(ws, lines):
    ws.column_dimensions["A"].width = 140
    for i, l in enumerate(lines, 1):
        ws.cell(row=i, column=1, value=l).alignment = Alignment(wrap_text=True, vertical="top")
        if l.startswith("# "): ws.cell(row=i, column=1).font = Font(bold=True, size=13)

def build(ctx, path):
    cfg, prov = ctx["cfg"], ctx["prov"]
    wb = Workbook(); ws = wb.active; ws.title = "README"
    em = cfg["EMERGENCE"]
    readme = [
        "# YC batch composition using YC's native taxonomy",
        f"Analysis window: batches with year >= {cfg['START_YEAR']} (END_YEAR = {cfg['END_YEAR'] or 'latest'}). Companies selected: {len(ctx['df'])} of {len(ctx['df_all'])} in the source directory.",
        f"Primary source: {prov['primary_source']['repo']} commit {prov['primary_source']['commit_sha']} ({prov['primary_source']['commit_date']}); source meta.json last_updated {prov['primary_source']['meta_last_updated']}; retrieved {prov['retrieval_date_utc']}.",
        "Coverage caveat: the directory lists publicly launched companies with YC profiles, not every company ever accepted into YC.",
        f"Secondary source (enrichment only): {prov['secondary_source']['repo']} commit {prov['secondary_source']['commit_sha']} ({prov['secondary_source']['commit_date']}).",
        f"Historical snapshots of companies/all.json (taxonomy change detection): {len(prov['historical_snapshots'])} git commits from {prov['historical_snapshots'][0][1][:10]} to {prov['historical_snapshots'][-1][1][:10]}.",
        "RFS sources: https://www.ycombinator.com/rfs (live) + Wayback Machine monthly captures 2014-2026 + ycombinator.com/rfs1..10.html essays (2009-2013).",
        "",
        "# Definitions",
        "industry = YC's single top-level `industry` field. subindustry = YC's `subindustry` string ('Parent -> Child'). tags = YC's `tags` array. Nothing is merged or renamed; derived fields are prefixed n_.",
        "share_pct = companies with the label / ALL companies in the batch (companies with zero tags are in the denominator). yc_tag_by_batch_share_of_tagged.csv gives the tagged-only denominator.",
        f"Partial batch = fewer than {cfg['MIN_FULL_BATCH']} companies (in progress): shown in matrices, excluded from rankings, 'latest' shares and emergence rules.",
        f"Low tag coverage batch = more than {cfg['TAG_COVERAGE_MAX_ZERO_TAG_PCT']}% of companies have zero tags: {', '.join(ctx['low_tag_batches']) or 'none'}. Excluded from TAG statistics only (kept in matrices).",
        f"'Latest' = latest full batch: {ctx['results']['industry']['stats']['latest_full_batch'].iloc[0]} (industry/subindustry) and {ctx['results']['tag']['stats']['latest_full_batch'].iloc[0]} (tags). Earliest-year share = pooled share over all {cfg['START_YEAR']} batches.",
        "",
        "# Emergence rules (applied to share-of-batch series over full batches, chronological)",
        f"early = mean of first {em['n_early_batches']} batches; late = mean of last {em['n_late_batches']} full batches.",
        f"Emerging: late - early >= {em['emerging_min_pp_gain']} pp AND late >= {em['emerging_min_ratio']} x early AND latest count >= {em['emerging_min_latest_count']}.",
        f"Declining: early >= {em['declining_min_early_share_pct']}% AND late <= {em['declining_max_ratio']} x early AND early - late >= {em['declining_min_pp_loss']} pp.",
        f"Spiky: peak >= {em['spiky_peak_over_median']} x median AND peak >= {em['spiky_min_peak_pct']}% (and not Emerging/Declining).",
        f"Persistent: mean >= {em['persistent_min_mean_pct']}% AND coefficient of variation <= {em['persistent_max_cv']}. Otherwise: Low-volume / mixed.",
        "",
        "# Co-occurrence",
        f"Periods: {', '.join(p['label'] for p in cfg['PERIODS'])}. Tags with >= {cfg['COOC_MIN_TAG_COMPANIES']} companies in the period; pairs with >= {cfg['COOC_MIN_PAIR_COUNT']} co-occurrences. lift = P(a,b)/(P(a)P(b)); Jaccard = |a&b|/|a|b|. strengthening/weakening = lift change >= +/-0.5 between first and last period.",
        "",
        "# RFS",
        "Requests for Startups parsed deterministically from HTML. 'RFS Label Mentions' is an exact, case-insensitive, word-boundary string match of YC label names in request title+body. It is NOT a semantic mapping.",
        "",
        "# Tabs", "Companies | Batch Summary | Industries | Industry Trends | Subindustries | Subindustry Trends | Tags | Tag Trends | Category Rankings | Tag Co-occurrence | Geography | Diversity | Taxonomy Changes | Data Quality | RFS | RFS Label Mentions",
    ]
    _text_sheet(ws, readme)
    R = ctx["results"]
    comp_cols = ["id", "name", "slug", "former_names", "url", "website", "batch", "n_batch_code", "n_year", "n_batch_is_partial", "one_liner", "long_description",
                 "team_size", "all_locations", "industry", "subindustry", "n_subindustry_child", "industries", "tags", "n_n_tags", "regions", "n_is_us", "n_is_remote",
                 "n_countries", "n_city", "stage", "status", "isHiring", "nonprofit", "top_company", "launched_at", "n_launched_date", "app_video_public",
                 "demo_day_video_public", "ycds_present", "ycds_year_founded", "ycds_city", "ycds_country", "ycds_has_postmortem", "ycds_linkedin_url"]
    dfe = ctx["df_enriched"]; comp_cols = [c for c in comp_cols if c in dfe.columns]
    _write(wb.create_sheet("Companies"), dfe[comp_cols], "Companies", link_cols=("url", "website", "ycds_linkedin_url"), max_width=45)
    _write(wb.create_sheet("Batch Summary"), ctx["bs"], "BatchSummary")
    _write(wb.create_sheet("Industries"), ctx["ind_inv"], "Industries")
    it = R["industry"]; _write(wb.create_sheet("Industry Trends"), it["stats"], "IndustryTrendStats")
    ws = wb["Industry Trends"]; start = len(it["stats"]) + 4
    ws.cell(row=start - 1, column=1, value="Share of batch (%) matrix").font = Font(bold=True)
    for r in [["industry"] + list(it["shares"].columns)] + [[i] + list(map(float, row)) for i, row in it["shares"].iterrows()]:
        ws.append(r)
    ws.append([]); ws.append(["Count matrix"])
    for r in [["industry"] + list(it["counts"].columns)] + [[i] + list(map(int, row)) for i, row in it["counts"].iterrows()]:
        ws.append(r)
    _write(wb.create_sheet("Subindustries"), ctx["sub_inv"], "Subindustries")
    st = R["subindustry"]; _write(wb.create_sheet("Subindustry Trends"), st["stats"], "SubindustryTrendStats")
    ws = wb["Subindustry Trends"]; ws.append([]); ws.append(["Share of batch (%) matrix"])
    for r in [["subindustry"] + list(st["shares"].columns)] + [[i] + list(map(float, row)) for i, row in st["shares"].iterrows()]:
        ws.append(r)
    _write(wb.create_sheet("Tags"), ctx["tag_inv"], "Tags")
    tt = R["tag"]; _write(wb.create_sheet("Tag Trends"), tt["stats"], "TagTrendStats")
    ws = wb["Tag Trends"]; ws.append([]); ws.append(["Share of batch (%) matrix (all companies in denominator)"])
    for r in [["tag"] + list(tt["shares"].columns)] + [[i] + list(map(float, row)) for i, row in tt["shares"].iterrows()]:
        ws.append(r)
    rk = pd.concat([R["industry"]["rank"].rename(columns={"industry": "category"}).assign(category_type="industry"),
                    R["subindustry"]["rank"].rename(columns={"subindustry": "category"}).assign(category_type="subindustry"),
                    R["tag"]["rank"].rename(columns={"tag": "category"}).assign(category_type="tag")])
    rk = rk[["category_type"] + [c for c in rk.columns if c != "category_type"]]
    _write(wb.create_sheet("Category Rankings"), rk, "CategoryRankings")
    ws = wb["Category Rankings"]; ws.append([]); ws.append(["Top-N per batch"])
    topn = pd.concat([R["industry"]["topn"].rename(columns={"industry": "category"}).assign(category_type="industry"), R["tag"]["topn"].rename(columns={"tag": "category"}).assign(category_type="tag")])
    ws.append(list(topn.columns))
    for row in topn.itertuples(index=False): ws.append([_san(v) for v in row])
    co = ctx["pairs"]; _write(wb.create_sheet("Tag Co-occurrence"), co, "TagCooccurrence")
    ws = wb["Tag Co-occurrence"]; ws.append([]); ws.append(["Pair change between first and last period"])
    pc = ctx["pair_change"]; ws.append(list(pc.columns))
    for row in pc.itertuples(index=False): ws.append([_san(None if (isinstance(v, float) and pd.isna(v)) else v) for v in row])
    _write(wb.create_sheet("Geography"), ctx["geo"], "Geography")
    ws = wb["Geography"]; ws.append([]); ws.append(["Countries by batch (from YC `regions`)"]); ws.append(list(ctx["countries"].columns))
    for row in ctx["countries"].itertuples(index=False): ws.append([_san(v) for v in row])
    ws.append([]); ws.append(["Macro-regions by batch (from YC `regions`)"]); ws.append(list(ctx["macro"].columns))
    for row in ctx["macro"].itertuples(index=False): ws.append([_san(v) for v in row])
    ws.append([]); ws.append(["Cities by batch (first location in `all_locations`)"]); ws.append(list(ctx["cities"].columns))
    for row in ctx["cities"].sort_values("count", ascending=False).head(3000).itertuples(index=False): ws.append([_san(v) for v in row])
    _write(wb.create_sheet("Diversity"), ctx["div"], "Diversity")
    th = ctx["th"]
    tx = wb.create_sheet("Taxonomy Changes")
    if th:
        summ = th["summary"]
        _text_sheet(tx, ["# Taxonomy change detection (snapshots of companies/all.json from yc-oss/api git history)",
                         f"{summ['n_snapshots']} snapshots, {summ['first_snapshot']} .. {summ['latest_snapshot']}; relabel events between consecutive snapshots: {summ['relabel_events']}.",
                         f"Tags first appearing after the first snapshot: {summ['tags_introduced_after_first_snapshot']}; tags absent in latest snapshot: {summ['tags_absent_in_latest']}; subindustries introduced after first snapshot: {summ['subindustries_introduced_after_first_snapshot']}.",
                         "Tables below: (1) retroactive relabeling by batch, (2) near-duplicate labels (listed, never merged), (3) tag relabel flows removed->added, (4) label vocabulary per snapshot."])
        r0 = 7
        for title, d in (("Retroactive relabeling by batch", th["retro"]), ("Near-duplicate labels", th["near_duplicates"]), ("Tag relabel flows", th["flows"]), ("Vocabulary per snapshot", th["vocab"])):
            tx.cell(row=r0, column=1, value=title).font = Font(bold=True); r0 += 1
            tx.append([]) if False else None
            for j, c in enumerate(d.columns, 1): tx.cell(row=r0, column=j, value=str(c)).font = Font(bold=True)
            r0 += 1
            for row in d.itertuples(index=False):
                for j, v in enumerate(row, 1): tx.cell(row=r0, column=j, value=_san(None if (isinstance(v, float) and pd.isna(v)) else v))
                r0 += 1
            r0 += 2
    dq = wb.create_sheet("Data Quality"); _write(dq, ctx["dq"], "DataQuality", max_width=80)
    dq.append([]); dq.append(["Source cross-check (our counts vs yc-oss/api list files and meta.json)"]); dq.append(list(ctx["xc"].columns))
    for row in ctx["xc"].itertuples(index=False): dq.append([_san(v) for v in row])
    dq.append([]); dq.append(["Secondary-source coverage by batch"]); dq.append(list(ctx["cov"].columns))
    for row in ctx["cov"].itertuples(index=False): dq.append([_san(v) for v in row])
    dq.append([]); dq.append([f"Spot check: {ctx['sc']['company_id'].nunique()} random companies x {ctx['sc']['field'].nunique()} fields compared with raw JSON; mismatches = {int((~ctx['sc']['match']).sum())}"])
    dq.append(list(ctx["sc"].columns))
    for row in ctx["sc"].itertuples(index=False): dq.append([_san(v) for v in row])
    _write(wb.create_sheet("RFS"), ctx["req"], "RFS", link_cols=("source_url",), max_width=50)
    ws = wb["RFS"]; ws.append([]); ws.append(["Editions"]); ws.append(list(ctx["ed"].columns))
    for row in ctx["ed"].itertuples(index=False): ws.append([_san(v) for v in row])
    ws.append([]); ws.append(["Title persistence across editions (exact title match)"]); ws.append(list(ctx["persist"].columns))
    for row in ctx["persist"].itertuples(index=False): ws.append([_san(v) for v in row])
    _write(wb.create_sheet("RFS Label Mentions"), ctx["ba"], "RFSLabelMentions")
    wb.save(path); return path
